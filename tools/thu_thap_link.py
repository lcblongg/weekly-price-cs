"""Bot 1 — thu thập link sản phẩm Apple của 5 kênh từ inputs/data.xlsx, xuất Excel tổng hợp để chọn link.

Chạy discovery ở chế độ dry-run (không ghi Supabase, không gửi Telegram), phạm vi theo config/scope.json.
Đầu ra: outputs/Tong hop link.xlsx — mỗi dòng 1 link đúng model + dung lượng + màu của một kênh.
Người dùng chép Model / Màu / Link cần lấy giá sang inputs/chon-link.xlsx rồi chạy bot lấy giá.
"""
import argparse
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from chains import CHAINS  # noqa: E402
from common import TZ  # noqa: E402

LABELS = {'tgdd': 'MW', 'cellphones': 'CPS', 'fpt': 'FPT', 'viettel': 'Viettel', 'phongvu': 'Phong Vũ'}
CHOICE = ROOT / 'inputs/chon-link.xlsx'
CHOICE_HEADER = ['Model', 'Màu', 'Link', 'Ghi chú']


def pinned_link(slug, config):
    """Link giữ đúng biến thể: MW cần ?code=, FPT/Phong Vũ đã có ?sku=; CPS/Viettel mỗi trang là một dung lượng."""
    url = config['url']
    if slug == 'tgdd' and config.get('variant_id'):
        parts = urlsplit(url)
        query = dict(parse_qsl(parts.query))
        query['code'] = str(config['variant_id'])
        url = urlunsplit(parts._replace(query=urlencode(query)))
    return url


def describe_row(slug, row):
    from identity import storage_of
    from product_standard import canonical_model
    config = row['config']
    name = config.get('source_product_name') or config.get('discovered_name') or config.get('product_name') or ''
    model = config.get('model_name') or canonical_model(name) or ''
    storage = config.get('storage') or storage_of(name) or ''
    return {'Kênh': LABELS[slug], 'Danh mục': config.get('category', ''), 'Model': model, 'Dung lượng': storage,
            'Màu': config.get('color') or '', 'Tên trên website': name, 'Link': pinned_link(slug, {**config, 'url': row['source_url']}),
            'Tình trạng': 'Đọc được giá' if row['status'] == 'ready' else 'Cần kiểm tra: ' + (row.get('reason') or '')}


def write_choice_template():
    if CHOICE.exists():
        return False
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    sheet = book.active
    sheet.title = 'Chọn link'
    sheet.append(CHOICE_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip('ABCD', (28, 22, 90, 30)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    CHOICE.parent.mkdir(parents=True, exist_ok=True)
    book.save(CHOICE)
    return True


def export(rows, path, started):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    order = {'Điện thoại': 0, 'Máy tính bảng': 1, 'Máy tính xách tay': 2, 'Đồng hồ thông minh': 3, 'Airpods': 4, 'AirPods': 4}
    rows = sorted(rows, key=lambda r: (order.get(r['Danh mục'], 9), r['Model'], r['Dung lượng'], r['Kênh'], r['Màu']))
    book = Workbook()
    sheet = book.active
    sheet.title = 'Tổng hợp link'
    header = list(rows[0]) if rows else ['Kênh', 'Danh mục', 'Model', 'Dung lượng', 'Màu', 'Tên trên website', 'Link', 'Tình trạng']
    sheet.append(header)
    for row in rows:
        sheet.append([row[h] for h in header])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F5F4A')
    for column, width in zip('ABCDEFGH', (10, 18, 26, 12, 20, 50, 80, 30)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    guide = book.create_sheet('Hướng dẫn')
    for line in [
        f'Thu thập lúc {started:%H:%M %d/%m/%Y}. Mỗi dòng = 1 link đúng model + dung lượng + màu của một kênh.',
        'Cách chọn: chép 3 cột Model, Màu, Link của các dòng cần lấy giá sang file inputs/chon-link.xlsx.',
        'Mỗi model chỉ theo dõi 1 màu; tên màu khác nhau giữa các kênh vẫn được (bot coi là cùng màu của model đó).',
        'Sau đó bấm "Lay gia.command" để bot vào đúng các link đã chọn, đẩy giá lên web và gửi Telegram nhóm.',
        'Dòng "Cần kiểm tra": website chưa hiện giá lúc thu thập (hết hàng, sắp về...). Vẫn có thể chọn.']:
        guide.append([line])
    guide.column_dimensions['A'].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', default=str(ROOT / 'outputs/Tong hop link.xlsx'))
    parser.add_argument('--work', default=str(ROOT / 'out/thu-thap-link'))
    parser.add_argument('--chain', default='all')
    args = parser.parse_args()
    started = datetime.now(TZ)
    work = Path(args.work)
    print('Đang thu thập link Apple của 5 kênh (có thể mất 20–40 phút; không ghi web, không gửi Telegram)…', flush=True)
    code = subprocess.run([sys.executable, 'discover_products.py', '--chain', args.chain, '--out', str(work), '--dry-run'], cwd=ROOT).returncode
    rows, notes = [], []
    for slug in CHAINS if args.chain == 'all' else [s.strip() for s in args.chain.split(',')]:
        folder = work / slug
        catalog = folder / 'catalog.json'
        if not catalog.exists():
            catalog = folder / 'catalog.failed.json'
        if not catalog.exists():
            notes.append(f'{LABELS.get(slug, slug)}: không thu thập được link')
            continue
        data = json.loads(catalog.read_text(encoding='utf-8'))
        if catalog.name == 'catalog.failed.json':
            notes.append(f'{LABELS[slug]}: có nguồn lỗi; danh sách có thể thiếu link')
        found = [describe_row(slug, r) for r in data if (r['config'].get('brand') or '').casefold() == 'apple']
        rows += found
        print(f'{LABELS[slug]}: {len(found)} link', flush=True)
    if not rows:
        print('Không thu thập được link nào; xem log trong ' + str(work), flush=True)
        return 1
    export(rows, Path(args.out), started)
    created = write_choice_template()
    print(f'Đã xuất {len(rows)} link → {args.out}', flush=True)
    if created:
        print(f'Đã tạo file chọn link trống: {CHOICE}', flush=True)
    for note in notes:
        print('Lưu ý: ' + note, flush=True)
    return 0 if code == 0 else 2


if __name__ == '__main__':
    sys.exit(main())
