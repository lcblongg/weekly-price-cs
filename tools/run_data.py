"""Một file data.xlsx có Tên sản phẩm / Màu / MW / CPS / FPT / Viettel / Phong Vũ."""
import argparse
import fcntl
import hashlib
import json
import os
from pathlib import Path
import sys
from datetime import datetime
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from apple_rules import RuleError, key, normalize_url, validate_products, plan
from common import database, PipelineError, TZ
from tools.update_web_local import run_job, wait_for_idle
from openpyxl import load_workbook

HEADERS={'mw':'tgdd','tgdd':'tgdd','cps':'cellphones','cellphones':'cellphones','fpt':'fpt','viettel':'viettel','vt':'viettel','phong vũ':'phongvu','pv':'phongvu'}

def parse_rows(rows):
    if not rows:raise RuleError('File data trống')
    header=[key(v or '') for v in rows[0]]
    model_col=next((i for i,v in enumerate(header) if v in ('tên sản phẩm','model')),None)
    if model_col is None or 'màu' not in header:raise RuleError('Cần cột Tên sản phẩm và Màu')
    color_col=header.index('màu');columns={}
    for i,h in enumerate(header):
        if h in HEADERS:
            slug=HEADERS[h]
            if slug in columns:raise RuleError('Trùng cột link của một kênh')
            columns[slug]=i
    if len(columns)!=5:raise RuleError('Cần 5 cột link MW, CPS, FPT, Viettel, Phong Vũ; link Hoàng Hà không dùng cho 5 kênh này')
    products={}
    for n,row in enumerate(rows[1:],2):
        def value(i):return row[i] if i<len(row) else None
        if not any(value(i) is not None for i in [model_col,color_col,*columns.values()]):continue
        for i in [model_col,color_col,*columns.values()]:
            if isinstance(value(i),str) and value(i).startswith('='):raise RuleError(f'Dòng {n}: dùng văn bản/link, không dùng công thức')
        model=' '.join(str(value(model_col) or '').split());color=str(value(color_col) or '').strip() or None
        if not model:raise RuleError(f'Dòng {n}: thiếu tên sản phẩm')
        if not key(model).startswith(('iphone','ipad','macbook','apple watch','airpods','mac mini','mac studio','imac','earpods','studio display','apple tv','homepod','apple pencil','airtag','magic ')):
            raise RuleError(f'Dòng {n}: chỉ nhận sản phẩm Apple')
        old=products.get(key(model))
        if old and key(old['color'] or '')!=key(color or ''):raise RuleError(f'Dòng {n}: cùng model có màu khác; không tự coi hai màu là tương đương')
        rule=products.setdefault(key(model),{'model':model,'color':color,'aliases':[],'urls':{}})
        for channel,i in columns.items():
            for text in str(value(i) or '').splitlines():
                if not text.strip():continue
                try:url=normalize_url(channel,text.strip())
                except (RuleError,ValueError):raise RuleError(f'Dòng {n}, cột {header[i]}: link sai kênh hoặc định dạng') from None
                urls=rule['urls'].setdefault(channel,[])
                if url not in urls:urls.append(url)
    if not products:raise RuleError('File chưa có sản phẩm')
    # Dùng cùng quy tắc với trang quản trị, kể cả giới hạn số model/link.
    return validate_products(list(products.values()))

def read_products(path):
    book=load_workbook(path,read_only=False,data_only=False,keep_links=False)
    try:
        sheet=book.worksheets[0]
        rows=[]
        for row in sheet.iter_rows():
            rows.append([c.value if c.data_type=='f' else c.hyperlink.target if c.hyperlink else c.value for c in row])
        return parse_rows(rows)
    finally:book.close()

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--file',default=str(ROOT.parent/'data.xlsx'))
    p.add_argument('--check',action='store_true')
    args=p.parse_args();file=Path(args.file)
    fingerprint=hashlib.sha256(file.read_bytes()).hexdigest();products=read_products(file)
    channels=list(dict.fromkeys(c for r in products for c in r.get('urls',{})))
    print(f'{len(products)} model Apple · {sum(len(u) for r in products for u in r.get("urls",{}).values())} link nhập · {len(channels)} kênh có link',flush=True)
    missing=[r['model'] for r in products if not r.get('urls')]
    if missing:print('Chưa nhập link:',', '.join(missing),flush=True)
    if args.check:return 0
    if not channels:raise RuleError('Chưa có link nào; điền link vào data.xlsx rồi chạy lại. Chưa đổi cấu hình hoặc giá.')
    db=database();lockpath=ROOT/'artifacts/locks/update-web-local.lock';lockpath.parent.mkdir(parents=True,exist_ok=True)
    with lockpath.open('a+') as lock:
        try:fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:raise PipelineError('Bot local đang chạy ở cửa sổ khác; không chạy trùng') from None
        wait_for_idle(db,channels,21600)
        if hashlib.sha256(file.read_bytes()).hexdigest()!=fingerprint:raise RuleError('File thay đổi trong lúc chờ; chạy lại để dùng bản mới')
        settings={r['key']:r for r in db.table('app_settings').select('key,value,revision').in_('key',['apple_colors','apple_models']).execute().data}
        # Chỉ giữ alias đã được xác nhận nếu màu yêu cầu không đổi.
        old={r['model']:r for r in settings['apple_colors']['value']['products']}
        for r in products:
            before=old.get(r['model'],{})
            if key(before.get('color') or '')==key(r['color'] or ''):r['aliases']=before.get('aliases',[])
        plan(products,settings['apple_models']['value'],current=settings['apple_colors']['value']['products'])
        backup=ROOT/f'out/data-runs/{datetime.now(TZ):%Y%m%d-%H%M%S}'
        backup.mkdir(parents=True,exist_ok=True)
        (backup/'configuration-before.json').write_text(json.dumps(settings,ensure_ascii=False))
        (backup/'input.json').write_text(json.dumps({'sha256':fingerprint,'products':products},ensure_ascii=False))
        run_job(db,'config_update',channels,21600,{'products':products,'renames':[],'expected_revision':settings['apple_colors']['revision'],'auto_followup':False})
        job=run_job(db,'daily_prices',channels,21600,{'selected_links':True})
        print('Đã kết thúc lượt đọc data.xlsx → cập nhật web. Không chụp ảnh, không gửi Telegram.',flush=True)
        return 2 if job['status']=='partial' else 0

if __name__=='__main__':
    try:sys.exit(main())
    except (PipelineError,RuleError,OSError) as exc:print('Chưa hoàn tất:',str(exc),flush=True);sys.exit(1)
    except Exception as exc:print('Chưa hoàn tất:',type(exc).__name__,'— kiểm tra kết nối/cấu hình.',flush=True);sys.exit(1)
