"""Bot 2 — lấy giá đúng các link trong inputs/chon-link.xlsx (cột Model, Màu, Link), đẩy lên web và gửi Telegram nhóm.

1. Đọc file chọn, nhận kênh theo domain, chuẩn hóa link (sai domain/định dạng → dừng, không chạy dở).
2. Ghép thành danh sách "Sản phẩm theo dõi": mỗi model 1 màu; tên màu khác nhau giữa các kênh là tên tương đương.
3. Sao lưu quy chuẩn đang có trên Supabase vào outputs/, rồi áp dụng danh sách mới (web chỉ hiện các model này).
4. Chạy bot giá ở chế độ chỉ-link-đã-chọn cho 5 kênh, lưu Supabase, gửi báo cáo Telegram (trừ khi --khong-gui-telegram).
"""
import argparse
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from apple_rules import URL_CHANNELS, RuleError, key, normalize_url, plan  # noqa: E402
from common import TZ, database  # noqa: E402

CHOICE = ROOT / 'inputs/chon-link.xlsx'


def channel_of(url):
    host = (urlsplit(url).hostname or '').lower()
    return next((slug for slug, spec in URL_CHANNELS.items() if host in spec['hosts']), None)


def read_choices(path):
    """[(dòng, model, màu, link)] — bỏ dòng trống. Cột nhận theo tiêu đề Model / Màu / Link."""
    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise RuleError('File chọn link trống')
    header = [str(c or '').strip().casefold() for c in rows[0]]
    try:
        cols = [header.index(name) for name in ('model', 'màu', 'link')]
    except ValueError:
        raise RuleError('File chọn link cần 3 cột tiêu đề: Model, Màu, Link') from None
    found = []
    for number, row in enumerate(rows[1:], 2):
        values = [' '.join(str(row[c] or '').split()) if c < len(row) else '' for c in cols]
        if any(values):
            found.append((number, *values))
    return found


def build_products(choices):
    """Danh sách quy chuẩn theo thứ tự file; lỗi gom đủ rồi báo một lần."""
    products, errors = {}, []
    for number, model, color, link in choices:
        if not model or not color or not link:
            errors.append(f'Dòng {number}: cần đủ Model, Màu, Link')
            continue
        slug = channel_of(link)
        if not slug:
            errors.append(f'Dòng {number}: link không thuộc 5 kênh (MW, CPS, FPT, Viettel, Phong Vũ)')
            continue
        try:
            url = normalize_url(slug, link)
        except RuleError as exc:
            errors.append(f'Dòng {number}: {exc}')
            continue
        rule = products.setdefault(key(model), {'model': model, 'color': color, 'aliases': [], 'urls': {}})
        if key(color) != key(rule['color']) and color not in rule['aliases']:
            rule['aliases'].append(color)  # cùng model, kênh gọi tên màu khác
        urls = rule['urls'].setdefault(slug, [])
        if url not in urls:
            urls.append(url)
    if errors:
        raise RuleError('\n'.join(errors))
    if not products:
        raise RuleError('File chọn link chưa có dòng nào')
    return list(products.values())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--file', default=str(CHOICE))
    parser.add_argument('--khong-gui-telegram', action='store_true', help='Chạy thử: lưu web nhưng không gửi nhóm')
    parser.add_argument('--chi-kiem-tra', action='store_true', help='Chỉ đọc và kiểm tra file chọn, không đổi gì')
    args = parser.parse_args()
    try:
        products = build_products(read_choices(Path(args.file)))
    except (RuleError, OSError) as exc:
        print('Không chạy được — sửa file chọn link rồi bấm lại:\n' + str(exc), flush=True)
        return 2
    for p in products:
        aliases = f" (tên khác: {', '.join(p['aliases'])})" if p['aliases'] else ''
        print(f"• {p['model']} — {p['color']}{aliases}: " + ', '.join(f"{URL_CHANNELS[s]['label']} {len(u)} link" for s, u in p['urls'].items()), flush=True)
    if args.chi_kiem_tra:
        return 0
    db = database()
    settings = {r['key']: r for r in db.table('app_settings').select('key,value,revision').in_('key', ['apple_colors', 'apple_models']).execute().data}
    try:
        colors, models = plan(products, settings['apple_models']['value'], [], settings['apple_colors']['value']['products'])
    except RuleError as exc:
        print('Danh sách không hợp lệ: ' + str(exc), flush=True)
        return 2
    changed = colors != settings['apple_colors']['value']
    if changed:
        backup = ROOT / f"outputs/sao-luu-quy-chuan-{datetime.now(TZ):%Y%m%d-%H%M%S}.json"
        backup.parent.mkdir(parents=True, exist_ok=True)
        backup.write_text(json.dumps({k: v['value'] for k, v in settings.items()}, ensure_ascii=False, indent=1), encoding='utf-8')
        db.rpc('set_app_configuration', {'p_colors': colors, 'p_models': models, 'p_revision': settings['apple_colors']['revision']}).execute()
        print(f'Đã cập nhật danh sách Sản phẩm theo dõi ({len(products)} model). Bản cũ lưu ở {backup.name}', flush=True)
    else:
        print('Danh sách Sản phẩm theo dõi không đổi', flush=True)
    command = [sys.executable, 'tools/cloud_worker.py', '--kind', 'daily_prices', '--chains', 'all', '--selected-links']
    if not args.khong_gui_telegram:
        command.append('--send-report')
    print('Đang lấy giá 5 kênh theo đúng link đã chọn…', flush=True)
    return subprocess.run(command, cwd=ROOT).returncode


if __name__ == '__main__':
    sys.exit(main())
