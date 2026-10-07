"""Kiểm thử lựa chọn model, persistence, Excel và mobile trên server port 3001."""
import asyncio,json,os
from pathlib import Path
from playwright.async_api import async_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
URL=os.environ.get('DASHBOARD_URL','http://127.0.0.1:3001')
KEY='weekly-price-cs:watchlist:v1'
async def main():
 async with async_playwright() as p:
  browser=await p.chromium.launch();page=await browser.new_page(viewport={'width':1440,'height':1000},accept_downloads=True)
  errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
  await page.goto(URL,wait_until='networkidle')
  await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).wait_for()
  await page.wait_for_function("localStorage.getItem('weekly-price-cs:watchlist:v1')!==null")
  initial=await page.locator('tbody tr').count();assert initial>0
  await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click()
  dialog=page.get_by_role('dialog');await dialog.get_by_role('textbox',name='Tìm model theo dõi').fill('iPhone 16')
  exact=dialog.get_by_role('checkbox',name='iPhone 16',exact=True);await exact.uncheck()
  await page.wait_for_function("JSON.parse(localStorage.getItem('weekly-price-cs:watchlist:v1')).hidden.length===1")
  await page.keyboard.press('Escape');await page.wait_for_timeout(300)
  count=await page.locator('tbody tr').count();assert 0<count<initial
  assert await page.locator('tbody').get_by_text('iPhone 16 Plus 256GB',exact=True).count() or await page.locator('tbody').get_by_text('iPhone 16 Plus 128GB',exact=True).count()
  await page.reload(wait_until='networkidle');await page.wait_for_timeout(300);assert await page.locator('tbody tr').count()==count
  # Kết hợp lựa chọn theo dõi với Category/kênh/model: model ẩn vẫn ẩn.
  demo=json.loads((ROOT/'web/data/demo.json').read_text())['rows']
  hidden_row=next(r for r in demo if r.get('model_name')=='iPhone 16' or r['product_name'].startswith('iPhone 16 128GB'))
  async def choose(name,value):
   await page.get_by_role('combobox',name=name).click();await page.get_by_role('option',name=value,exact=True).click()
  await choose('Chọn Category','Điện thoại');await choose('Chọn đại lý',hidden_row['chain_name']);await choose('Chọn model','iPhone 16');await page.wait_for_timeout(300)
  assert await page.locator('tbody tr').count()==0
  await choose('Chọn model','Tất cả model');await choose('Chọn đại lý','Tất cả đại lý');await choose('Chọn Category','Tất cả Category');await page.wait_for_timeout(300)
  assert await page.locator('tbody tr').count()==count
  async with page.expect_download() as download:await page.get_by_role('button',name='Xuất Excel',exact=True).click()
  path=ROOT/'artifacts/watchlist-export-test.xlsx';await (await download.value).save_as(str(path));book=load_workbook(path)
  assert book['Giá 7 ngày'].max_row==count+1
  assert not any('iPhone 16 128GB' in str(book['Giá 7 ngày'].cell(i,3).value) for i in range(2,book['Giá 7 ngày'].max_row+1));book.close()
  await page.get_by_role('button',name='Khuyến mãi đại lý',exact=True).click();assert await page.locator('.promo-list article').count()==count
  await page.get_by_role('button',name='Bảng giá thị trường',exact=True).click()
  await page.get_by_role('button',name='Theo ngày',exact=True).click();await page.wait_for_timeout(300);assert await page.locator('tbody tr').count()==count
  async with page.expect_download() as daily_download:await page.get_by_role('button',name='Xuất Excel',exact=True).click()
  day_path=ROOT/'artifacts/watchlist-day-export-test.xlsx';await (await daily_download.value).save_as(str(day_path));day_book=load_workbook(day_path)
  assert day_book['Giá theo ngày'].max_row==count+1
  expected={r['sku']:r for r in demo}
  for i in range(2,count+2):assert day_book['Giá theo ngày'].cell(i,5).value==expected[day_book['Giá theo ngày'].cell(i,4).value]['promo_price']
  day_book.close()
  await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click();await dialog.get_by_role('button',name='Bỏ chọn tất cả' ,exact=True).click();await page.keyboard.press('Escape');await page.wait_for_timeout(300);assert await page.locator('tbody tr').count()==0
  await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click();await dialog.get_by_role('button',name='Chọn tất cả',exact=True).click();await page.keyboard.press('Escape');await page.wait_for_timeout(300);assert await page.locator('tbody tr').count()==initial
  # Model mới mặc định bật, nhãn Mới tồn tại sau reload.
  state=await page.evaluate('(key)=>JSON.parse(localStorage.getItem(key))',KEY);entry=next(m for m in state['models'] if m['name']=='iPhone 16')
  state['models']=[m for m in state['models'] if m['id']!=entry['id']]
  await page.evaluate('([key,value])=>localStorage.setItem(key,JSON.stringify(value))',[KEY,state]);await page.reload(wait_until='networkidle');await page.wait_for_timeout(300)
  await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click();await dialog.get_by_role('textbox',name='Tìm model theo dõi').fill('iPhone 16')
  assert await dialog.get_by_role('checkbox',name='iPhone 16',exact=True).is_checked();assert await dialog.get_by_text('Mới',exact=True).count()>0
  await dialog.get_by_role('button',name='Khôi phục mặc định',exact=True).click();await page.keyboard.press('Escape')
  await page.set_viewport_size({'width':390,'height':844});await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click();await page.wait_for_timeout(300)
  assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth');box=await dialog.bounding_box();assert box and box['width']<=390
  await page.screenshot(path=str(ROOT/'artifacts/watchlist-mobile.png'));await page.keyboard.press('Escape');await page.set_viewport_size({'width':1440,'height':1000});await page.get_by_role('button',name='Sản phẩm theo dõi',exact=True).click();await page.screenshot(path=str(ROOT/'artifacts/watchlist-desktop.png'))
  assert not errors,errors
  print(json.dumps({'products':initial,'after_hide':count,'reload':True,'excel':True,'promos':True,'day':True,'new_model':True,'mobile':True,'page_errors':errors},ensure_ascii=False));await browser.close()
asyncio.run(main())
