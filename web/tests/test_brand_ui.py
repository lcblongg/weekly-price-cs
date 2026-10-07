"""Kiểm tra hãng + kênh/category/model, CTKM, Excel, catalog chưa có giá và mobile."""
import asyncio,json
from collections import Counter
from pathlib import Path
from playwright.async_api import async_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
async def main():
 data=json.loads((ROOT/'web/data/demo.json').read_text())
 source=json.loads((ROOT/Path(data['prices_source'])/'cellphones/prices.json').read_text())
 samsung=[r for r in source if r.get('brand')=='Samsung' and r['category']=='Điện thoại']
 model,count=Counter(r['model_name'] for r in samsung if r.get('model_name')).most_common(1)[0]
 async with async_playwright() as pw:
  browser=await pw.chromium.launch();page=await browser.new_page(viewport={'width':1440,'height':1000},accept_downloads=True)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  await page.goto('http://127.0.0.1:3000',wait_until='networkidle')
  async def choose(label,value):
   await page.get_by_role('combobox',name=label).click();await page.get_by_role('option',name=value,exact=True).click();await page.wait_for_timeout(200)
  await choose('Chọn đại lý','CellphoneS');await choose('Chọn hãng','Samsung');await choose('Chọn Category','Điện thoại');await choose('Chọn model',model)
  assert await page.locator('tbody tr').count()==count
  async with page.expect_download() as info:await page.get_by_role('button',name='Xuất Excel',exact=True).click()
  file=ROOT/'artifacts/brand-filter-export.xlsx';await (await info.value).save_as(str(file));book=load_workbook(file)
  assert book['Giá 7 ngày'].max_row==count+1
  assert all(book['Giá 7 ngày'].cell(i,2).value=='CellphoneS' for i in range(2,count+2));book.close()
  await page.get_by_role('button',name='Khuyến mãi đại lý',exact=True).click();assert await page.locator('.promo-list article').count()==count
  await page.get_by_role('button',name='Bảng giá thị trường',exact=True).click();await page.get_by_role('button',name='Theo ngày',exact=True).click();await page.wait_for_timeout(200);assert await page.locator('tbody tr').count()==count
  await page.get_by_role('button',name='Theo tuần',exact=True).click()
  await choose('Chọn hãng','Apple');assert await page.get_by_role('combobox',name='Chọn model').inner_text()=='Tất cả model'
  # Catalog có model chưa đọc giá: hiển thị trong dropdown + danh sách riêng, xuất không có giá giả.
  observed={r.get('model_name') for r in data['rows']+data['issues'] if r['chain_name']=='TGDD'}
  pending=next(c for c in data['catalog'] if c['chain_name']=='TGDD' and c['status']=='ready' and c.get('model_name') and c['model_name'] not in observed)
  await choose('Chọn đại lý','TGDD');await choose('Chọn hãng',pending['brand']);await choose('Chọn Category',pending['category']);await choose('Chọn model',pending['model_name'])
  assert await page.locator('tbody tr').count()==0
  await page.locator('.pending-catalog summary').click();assert await page.locator('.pending-catalog article').count()>0
  async with page.expect_download() as info:await page.get_by_role('button',name='Xuất Excel',exact=True).click()
  file=ROOT/'artifacts/catalog-pending-export.xlsx';await (await info.value).save_as(str(file));book=load_workbook(file)
  assert book['Giá 7 ngày'].max_row==1;assert book['Chưa có kết quả giá'].max_row>1;assert not any(c.data_type=='f' for ws in book for row in ws for c in row);book.close()
  await page.set_viewport_size({'width':390,'height':844});await page.wait_for_timeout(300);assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth')
  await page.screenshot(path=str(ROOT/'artifacts/brand-filter-mobile.png'));assert not errors,errors
  print(json.dumps({'chain':'CellphoneS','brand':'Samsung','category':'Điện thoại','model':model,'rows':count,'pending_model':pending['model_name'],'excel':True,'mobile':True,'page_errors':errors},ensure_ascii=False));await browser.close()
asyncio.run(main())
