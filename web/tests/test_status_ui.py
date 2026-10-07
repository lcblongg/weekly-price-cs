"""Đối chiếu trạng thái thật trên desktop/mobile và Excel; NULL không phải giá 0."""
import asyncio,json
from pathlib import Path
from openpyxl import load_workbook
from playwright.async_api import async_playwright
ROOT=Path(__file__).resolve().parents[2]
async def main():
 data=json.loads((ROOT/'web/data/demo.json').read_text())
 target=next(r for r in data['rows'] if r['chain_name']=='CellphoneS' and r['promo_price'] is None and r.get('model_name'))
 expected=[r for r in data['rows'] if r['chain_name']=='CellphoneS' and r.get('category')==target['category'] and r.get('brand')==target['brand'] and r.get('model_name')==target['model_name']]
 async with async_playwright() as pw:
  browser=await pw.chromium.launch();page=await browser.new_page(viewport={'width':1440,'height':1000},accept_downloads=True)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  await page.goto('http://127.0.0.1:3000',wait_until='networkidle')
  async def choose(label,value):
   await page.get_by_role('combobox',name=label).click();await page.get_by_role('option',name=value,exact=True).click();await page.wait_for_timeout(200)
  for label,value in [('Chọn đại lý','CellphoneS'),('Chọn hãng',target['brand']),('Chọn Category',target['category']),('Chọn model',target['model_name'])]:await choose(label,value)
  assert await page.locator('tbody tr').count()==len(expected)
  await page.get_by_role('button',name='Lịch sử 7 ngày',exact=True).click()
  assert await page.get_by_text('Có trạng thái · không hiện giá',exact=True).count()==sum(r['promo_price'] is None for r in expected)
  async with page.expect_download() as download:await page.get_by_role('button',name='Xuất Excel',exact=True).click()
  path=ROOT/'artifacts/status-export-test.xlsx';await (await download.value).save_as(str(path))
  book=load_workbook(path);sheet=book['Giá 7 ngày'];assert sheet.max_row==len(expected)+1
  by_sku={r['sku']:r for r in expected}
  for row in list(sheet.values)[1:]:
   ref=by_sku[row[3]];col=4+int(ref['business_date'][-2:])-5
   assert row[col]==ref['promo_price']
   if ref['promo_price'] is None:assert row[-1] and 'hết hàng' in row[-1].lower()
  assert 'Trạng thái 7 ngày' in book.sheetnames;book.close()
  await page.get_by_role('button',name='Theo ngày',exact=True).click();await page.wait_for_timeout(300)
  assert await page.get_by_text('Không hiện giá · xem trạng thái',exact=True).count()==sum(r['promo_price'] is None for r in expected)
  await page.set_viewport_size({'width':390,'height':844});await page.wait_for_timeout(200)
  assert await page.locator('.mobile-products article').count()==len(expected)
  assert await page.locator('.mobile-products .product-status').count()>0
  assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth')
  await page.screenshot(path=str(ROOT/'artifacts/status-mobile.png'));assert not errors,errors
  print(json.dumps({'model':target['model_name'],'records':len(expected),'status_only':sum(r['promo_price'] is None for r in expected),'desktop':True,'mobile':True,'excel':True},ensure_ascii=False));await browser.close()
asyncio.run(main())
