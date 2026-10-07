"""Kiểm tra dashboard đang chạy (npm run dev, chế độ demo) với dữ liệu trong web/data/demo.json.
Kỳ vọng được tính từ chính dữ liệu, không cố định số dòng. Chạy: .venv/bin/python web/tests/test_dashboard_ui.py
Kiểm: đủ 5 kênh, lọc kênh/category/model, tìm kiếm, theo ngày/tuần, 7 cột ngày với ngày thiếu là '—',
xuất Excel đúng dữ liệu đang lọc (tiền là số, ngày thiếu là ô trống, không công thức), tab Cần kiểm tra, mobile.
"""
import asyncio
import json
import os
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from playwright.async_api import async_playwright

ROOT = Path(__file__).resolve().parents[2]
DEMO = json.loads((ROOT / 'web/data/demo.json').read_text(encoding='utf-8'))
import os
URL = os.environ.get('DASHBOARD_URL', 'http://127.0.0.1:3000')
CHAINS = ['TGDD', 'CellphoneS', 'FPT Shop', 'Viettel Store', 'Phong Vũ']
CATEGORIES = ['Điện thoại', 'Máy tính bảng', 'Máy tính xách tay', 'Đồng hồ thông minh', 'AirPods']
MIN_PER_CHAIN = 10      # ngưỡng tối thiểu độc lập với dữ liệu: demo thiếu kênh/thiếu dòng phải làm test FAIL


def check_dataset():
    """Kiểm tra dữ liệu demo trước khi so với UI: thiếu kênh, thiếu nhóm hàng, giá không hợp lệ,
    nhiều ngày giả hoặc lệch với file cào nguồn đều là lỗi."""
    rows, issues = DEMO['rows'], DEMO.get('issues', [])
    assert rows, 'demo.json không có giá'
    days = {r['business_date'] for r in rows}
    stamps = {r['captured_at'] for r in rows}
    assert len(days) == 1, f'Demo phải là đúng một ngày cào thật, có {days}'
    per_chain_stamps = {(r['chain_name'], r['captured_at']) for r in rows}
    assert all(r['captured_at'] for r in rows), 'Mọi bản ghi phải có thời điểm đọc thật'
    # Một lượt bổ sung giữ thời điểm riêng của các giá đã đọc trước đó.
    per_chain = Counter(r['chain_name'] for r in rows)
    missing = [c for c in CHAINS if per_chain[c] < MIN_PER_CHAIN]
    assert not missing, f'Thiếu dữ liệu kênh: {missing} ({dict(per_chain)})'
    cats = Counter(r.get('category') for r in rows)
    assert all(cats[c] > 0 for c in CATEGORIES), f'Thiếu nhóm hàng: {dict(cats)}'
    assert all((isinstance(r['promo_price'], int) and r['promo_price'] > 0) or (r['promo_price'] is None and r['promo_text'].startswith('[Tình trạng] ')) for r in rows)
    assert all(r['original_price'] is None or (r['promo_price'] is not None and r['original_price'] >= r['promo_price']) for r in rows)
    assert len({(r['chain_name'], r['sku']) for r in rows}) == len(rows), 'SKU trùng trong một lượt'
    sources = sorted((ROOT / DEMO.get('prices_source','artifacts/full/prices')).glob('*/prices.json'))
    if sources:  # Đối chiếu với kết quả bot giá đã dùng để dựng demo: không thêm/bớt dòng.
        source_rows = sum(len(json.loads(p.read_text(encoding='utf-8'))) for p in sources)
        assert source_rows == len(rows), f'demo.json ({len(rows)}) lệch prices.json của bot ({source_rows})'
    return per_chain, len(issues)


def latest_keys(rows):
    return {(r['chain_name'], r['sku']) for r in rows}


async def count_rows(page):
    await page.wait_for_timeout(300)
    return await page.locator('tbody tr').count()


async def choose(page, combo, option):
    await page.get_by_role('combobox', name=combo).click()
    await page.get_by_role('option', name=option, exact=True).click()


async def main():
    check_dataset()
    rows = DEMO['rows']
    day = sorted({r['business_date'] for r in rows})[-1]
    week = sorted({r['week_start'] for r in rows})[-1]
    week_rows = [r for r in rows if r['week_start'] == week]
    chains = Counter(c for c, _ in latest_keys(week_rows))
    assert len(chains) == 5, f'Demo phải có đủ 5 kênh, hiện có {dict(chains)}'
    result = {'week_start': week, 'products': len(latest_keys(week_rows)), 'by_chain': dict(chains)}
    async with async_playwright() as pw:
        browser = await pw.chromium.launch()
        page = await browser.new_page(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        await page.goto(URL, wait_until='networkidle')
        await page.locator('tbody tr').first.wait_for()
        assert await count_rows(page) == result['products'], (await count_rows(page), result['products'])
        await page.screenshot(path=str(ROOT / 'artifacts/dashboard-desktop.png'))

        # Lọc từng kênh: số dòng = số SKU của kênh trong tuần.
        for chain, expected in chains.items():
            await choose(page, 'Chọn đại lý', chain)
            assert await count_rows(page) == expected, (chain, await count_rows(page), expected)
        await choose(page, 'Chọn đại lý', 'Tất cả đại lý')

        # Category + model áp dụng đồng thời; chọn category có nhiều SKU nhất.
        cats = Counter(r.get('category') for r in {(r['chain_name'], r['sku']): r for r in week_rows}.values())
        category = cats.most_common(1)[0][0]
        await choose(page, 'Chọn Category', category)
        assert await count_rows(page) == cats[category]
        models = Counter(r.get('model_name') for r in {(r['chain_name'], r['sku']): r for r in week_rows}.values()
                         if r.get('category') == category and r.get('model_name'))
        model, model_count = models.most_common(1)[0]
        await choose(page, 'Chọn model', model)
        assert await count_rows(page) == model_count, (model, await count_rows(page), model_count)
        result.update(category=category, category_rows=cats[category], model=model, model_rows=model_count)

        # Xuất Excel đúng các dòng đang lọc; tuần: 7 cột ngày, ngày thiếu để trống.
        async with page.expect_download() as info:
            await page.get_by_role('button', name='Xuất Excel', exact=True).click()
        export = ROOT / 'artifacts/dashboard-export-test.xlsx'
        await (await info.value).save_as(str(export))
        book = load_workbook(export)
        sheet = book['Giá 7 ngày']
        assert sheet.max_row == model_count + 1
        dates = [(date.fromisoformat(week) + timedelta(days=i)).isoformat() for i in range(7)]
        day_col = 5 + dates.index(day)
        for r in range(2, sheet.max_row + 1):
            expected_row=next(item for item in week_rows if item['sku']==sheet.cell(r,4).value and item['chain_name']==sheet.cell(r,2).value)
            assert sheet.cell(r,day_col).value==expected_row['promo_price']
            for c in range(5, 12):
                if c != day_col:
                    assert sheet.cell(r, c).value is None, 'Ngày chưa cào phải để trống'
        assert not any(cell.data_type == 'f' for ws in book for row in ws for cell in row)
        assert 'Cần kiểm tra' in book.sheetnames
        book.close()
        await choose(page, 'Chọn model', 'Tất cả model')
        await choose(page, 'Chọn Category', 'Tất cả Category')

        # Lịch sử 7 ngày: ngày chưa có dữ liệu hiện '—', không lấy giá ngày khác.
        await page.get_by_role('button', name='Lịch sử 7 ngày', exact=True).click()
        assert await page.locator('thead th').count() == 10
        first = page.locator('tbody tr').first
        cells = [await first.locator('td').nth(2 + i).inner_text() for i in range(7)]
        for i, text in enumerate(cells):
            if dates[i]==day:assert '₫' in text or 'Có trạng thái' in text,(dates[i],text)
            else:assert text=='—',(dates[i],text)
        await page.screenshot(path=str(ROOT / 'artifacts/dashboard-week.png'))

        # Theo ngày: đúng ngày có dữ liệu; ngày trước đó trống.
        await page.get_by_role('button', name='Theo ngày', exact=True).click()
        assert await count_rows(page) == len(latest_keys([r for r in rows if r['business_date'] == day]))
        prior = (date.fromisoformat(day) - timedelta(days=1)).isoformat()
        label = f'{prior[8:10]}/{prior[5:7]}/{prior[:4]}'
        await page.get_by_role('combobox', name='Chọn ngày').click()
        await page.get_by_role('option', name=label, exact=True).click()
        await page.get_by_role('heading', name='Chưa có dữ liệu ngày này').wait_for()
        await page.get_by_role('button', name='Theo tuần', exact=True).click()

        # Tab cần kiểm tra: số mục = số issue của lượt mới nhất.
        issues = {(i['chain_name'], i['source_url'], i['sku']) for i in DEMO.get('issues', []) if i['week_start'] == week}
        await page.get_by_role('button', name='Cần kiểm tra').first.click()
        if issues:
            await page.locator('.issue-list article').first.wait_for()
            assert await page.locator('.issue-list article').count() == len(issues)
        result['issues'] = len(issues)
        await page.get_by_role('button', name='Bảng giá thị trường', exact=True).click()

        await page.set_viewport_size({'width': 390, 'height': 844})
        await page.wait_for_timeout(500)
        await page.screenshot(path=str(ROOT / 'artifacts/dashboard-mobile.png'))
        assert await page.evaluate('document.documentElement.scrollWidth<=window.innerWidth')
        assert not errors, errors
        result.update(mobile_no_page_overflow=True, page_errors=errors)
        (ROOT / 'artifacts/dashboard-ui-check.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
        print('Đạt:', json.dumps(result, ensure_ascii=False))
        await browser.close()

asyncio.run(main())
