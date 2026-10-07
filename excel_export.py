"""Xuất file link tự động trên GitHub runner, nơi không có Artifact Tool của ứng dụng Codex.
Chỉ dữ liệu đã quét được; nguồn lỗi có sheet riêng. Không tạo sản phẩm giả khi quét lỗi.
"""
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from common import TZ


def export_discovery(rows, diagnostics, path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    timestamp = datetime.now(TZ).replace(tzinfo=None)
    products = workbook.create_sheet('Link sản phẩm')
    products.append(['Đơn vị', 'Danh mục nguồn', 'Hãng nguồn', 'Tên hiển thị', 'SKU',
                     'URL sản phẩm', 'Trạng thái', 'Lý do cần kiểm tra', 'Thời điểm tìm thấy', 'URL danh mục'])
    for row in rows:
        item = row['config']
        products.append([row['chain_name'], item.get('category', ''), item.get('brand', ''),
                         item.get('product_name') or item.get('discovered_name', ''), item.get('sku', ''),
                         row['source_url'], 'Đủ điều kiện cào' if row['status'] == 'ready' else 'Cần kiểm tra',
                         row['reason'], datetime.fromisoformat(item['discovered_at']).astimezone(TZ).replace(tzinfo=None) if item.get('discovered_at') else timestamp, '\n'.join(item.get('listing_urls', []))])
    sources = workbook.create_sheet('Nguồn quét')
    sources.append(['Đơn vị', 'Danh mục', 'Hãng', 'URL danh mục', 'Số link tìm thấy', 'Kết quả', 'Chi tiết'])
    for item in diagnostics:
        sources.append([item['chain_name'], item.get('category', ''), item.get('brand', ''),
                        item['url'], item.get('count', 0), item['status'], item.get('error', '')])
    for sheet, widths in [(products, [18, 24, 14, 48, 42, 70, 22, 60, 25, 70]),
                          (sources, [18, 24, 14, 70, 22, 24, 65])]:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
            cell.fill = PatternFill('solid', fgColor='183153')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        for i, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 48
            for cell in row:
                cell.font = Font(name='Arial', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                # Tên sản phẩm/ô nguồn là dữ liệu, không được trở thành công thức Excel.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if isinstance(cell.value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix('.tmp.xlsx')
    workbook.save(temporary)
    workbook.close()
    temporary.replace(output)
