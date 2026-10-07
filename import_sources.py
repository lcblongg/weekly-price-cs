"""Đọc workbook người dùng; dữ liệu ô không được thực thi như lệnh/công thức.
File gồm sheet tên đơn vị, cột Danh mục / Hãng / Link. Giữ query lọc của URL nguồn.
"""
import argparse
import copy
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
from common import PipelineError
from scraper import validate_url

SHEETS = {'MWG': 'TGDD', 'CPS': 'CellphoneS', 'FPT': 'FPT Shop',
          'VT': 'Viettel Store', 'PV': 'Phong Vũ'}
HEADERS = ('Danh mục', 'Hãng', 'Link')


def read_sources(path):
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=False)
    rows, seen = [], set()
    try:
        for sheet in workbook:
            populated = [row for row in sheet.iter_rows() if any(c.value is not None or c.hyperlink for c in row)]
            if not populated:
                continue
            chain = SHEETS.get(sheet.title.strip().upper())
            if not chain:
                raise PipelineError(f'Sheet chưa được ánh xạ đại lý: {sheet.title}')
            header = {str(c.value).strip().casefold(): c.column - 1 for c in populated[0] if c.value is not None}
            if any(name.casefold() not in header for name in HEADERS):
                raise PipelineError(f'{sheet.title}: cần cột Danh mục, Hãng, Link')
            for row in populated[1:]:
                selected = [row[header[name.casefold()]] for name in HEADERS]
                if any(c.data_type == 'f' for c in selected):
                    raise PipelineError(f'{sheet.title} dòng {row[0].row}: dùng văn bản hoặc hyperlink, không dùng công thức')
                category, brand = [str(c.value or '').strip() for c in selected[:2]]
                link_cell = selected[2]
                url = str((link_cell.hyperlink.target if link_cell.hyperlink else link_cell.value) or '').strip()
                if not category or not brand or not url:
                    raise PipelineError(f'{sheet.title} dòng {row[0].row}: thiếu danh mục, hãng hoặc link')
                validate_url(chain, url)
                key = (chain, category, brand, url)
                if key in seen:
                    continue
                seen.add(key)
                rows.append(dict(chain_name=chain, category=category, brand=brand, url=url,
                                 input_sheet=sheet.title, input_row=row[0].row))
    finally:
        workbook.close()
    if not rows:
        raise PipelineError('Workbook không có nguồn quét')
    return rows


def build_sources(rows, adapters):
    result = []
    # Loại URL danh mục đã biết ở tất cả sheet khỏi tập URL sản phẩm.
    seed_urls = list({row['url'] for row in rows})
    for row in rows:
        if row['chain_name'] not in adapters:
            raise PipelineError('Thiếu adapter cho đại lý trong workbook')
        source = copy.deepcopy(adapters[row['chain_name']])
        patterns = source.pop('category_url_patterns')
        if row['category'] not in patterns:
            raise PipelineError(f'Chưa có quy tắc link cho danh mục {row["category"]}')
        # Workbook quyết định nguồn; cấu hình adapter không thay thế danh sách người dùng.
        source.update(chain_name=row['chain_name'], seeds=[row['url']],
                      product_url_pattern=patterns[row['category']], category=row['category'],
                      brand=row['brand'], input_sheet=row['input_sheet'], input_row=row['input_row'],
                      excluded_urls=seed_urls)
        source['identity_family'] = 'iphone' if row['category'] == 'Điện thoại' and row['brand'] == 'Apple' else 'unsupported'
        result.append(source)
    return result


def prepare(path, adapter_path, output_dir='artifacts'):
    rows = read_sources(path)
    adapters = json.loads(Path(adapter_path).read_text(encoding='utf-8'))
    sources = build_sources(rows, adapters)
    if output_dir is None:
        return sources
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    (output / 'sources.json').write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'input-summary.json').write_text(json.dumps({
        'file': Path(path).name, 'total': len(rows),
        'chains': dict(Counter(r['chain_name'] for r in rows)),
        'categories': dict(Counter(r['category'] for r in rows)), 'rows': rows,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    return sources


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', default='inputs/data.xlsx')
    parser.add_argument('--adapters', default='config/retailer_adapters.json')
    args = parser.parse_args()
    try:
        sources = prepare(args.input, args.adapters)
        print(f'Đã đọc {len(sources)} nguồn từ Excel, tạo artifacts/sources.json')
    except Exception as exc:
        print(str(exc) if isinstance(exc, PipelineError) else type(exc).__name__)
        raise SystemExit(1)
